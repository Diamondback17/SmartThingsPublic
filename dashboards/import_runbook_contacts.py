#!/usr/bin/env python3
"""One-off/repeatable syncer: fills in primary_phone / escalation_phone on
existing hyperview_runbook entries by matching primary_name / escalation_name
against a Contact Name -> Contact Phone directory (e.g. Book2.xlsx).

Usage:
    python3 import_runbook_contacts.py <path-to-Book2.xlsx> [--db path/to/runbook.db] [--set-by name]

Safe to re-run: matches are by name (case-insensitive), and a matched
contact's phone number is synced from the directory every time - if a
number changes in the spreadsheet, re-running this updates every runbook
entry that uses that contact. Runbook contact names with no match in the
directory are left untouched and reported, so you know which ones still
need a number added to the source sheet. Expects a single sheet with header
row:
    Contact Name, Contact Phone
"""
import argparse
import os
import sys

import openpyxl


def _clean(value):
    if value is None:
        return ""
    return str(value).strip()


def load_directory(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("Spreadsheet is empty")
    directory, display_names = {}, {}
    for row in rows[1:]:
        if row is None or all(c is None for c in row):
            continue
        name = _clean(row[0]) if len(row) > 0 else ""
        phone = _clean(row[1]) if len(row) > 1 else ""
        if not name or not phone:
            continue
        key = name.lower()
        directory[key] = phone
        display_names[key] = name
    return directory, display_names


def compute_changes(entries, directory):
    used_names, unmatched_names, changes = set(), set(), []
    for e in entries:
        new_primary_phone = e["primary_phone"]
        new_escalation_phone = e["escalation_phone"]
        changed = False
        if e["primary_name"]:
            key = e["primary_name"].strip().lower()
            if key in directory:
                used_names.add(key)
                if directory[key] != e["primary_phone"]:
                    new_primary_phone = directory[key]
                    changed = True
            else:
                unmatched_names.add(e["primary_name"].strip())
        if e["escalation_name"]:
            key = e["escalation_name"].strip().lower()
            if key in directory:
                used_names.add(key)
                if directory[key] != e["escalation_phone"]:
                    new_escalation_phone = directory[key]
                    changed = True
            else:
                unmatched_names.add(e["escalation_name"].strip())
        if changed:
            changes.append((e, new_primary_phone, new_escalation_phone))
    return changes, used_names, unmatched_names


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx_path", help="Path to the contact directory (e.g. Book2.xlsx)")
    parser.add_argument("--db", default=None, help="Path to runbook.db (defaults to RUNBOOK_DB_PATH env var, or ./runbook.db)")
    parser.add_argument("--set-by", default="contact-sync", help="Value recorded as updated_by for updated rows (default: 'contact-sync')")
    parser.add_argument("--dry-run", action="store_true", help="Report what would change, without writing to the db")
    args = parser.parse_args()

    if args.db:
        os.environ["RUNBOOK_DB_PATH"] = args.db

    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import bridge_portal as bp

    directory, display_names = load_directory(args.xlsx_path)
    print(f"Loaded {len(directory)} contact(s) from {args.xlsx_path}")

    entries = bp._hyperview_runbook_entries()
    changes, used_names, unmatched_names = compute_changes(entries, directory)
    print(f"{len(changes)} of {len(entries)} runbook entries need a phone number added/updated")

    if unmatched_names:
        print(f"\n{len(unmatched_names)} contact name(s) used in the runbook have no match in the directory:")
        for n in sorted(unmatched_names):
            print(f"  - {n}")

    unused = sorted(set(directory) - used_names)
    if unused:
        print(f"\n{len(unused)} contact(s) in the directory aren't used by any current runbook entry (informational only):")
        for key in unused:
            print(f"  - {display_names[key]}")

    if args.dry_run:
        print("\n--dry-run: not writing to the database.")
        return

    if not changes:
        print("\nNothing to update.")
        return

    db_path = os.environ.get("RUNBOOK_DB_PATH", getattr(bp, "RUNBOOK_DB_PATH", "runbook.db"))
    print(f"\nWriting to {db_path} ...")
    for e, primary_phone, escalation_phone in changes:
        fields = {c: e[c] for c in bp.HYPERVIEW_RUNBOOK_FIELDS}
        fields["primary_phone"] = primary_phone
        fields["escalation_phone"] = escalation_phone
        bp._save_hyperview_runbook_entry(e["id"], fields, args.set_by, ignore_incomplete=bool(e["ignore_incomplete"]))
    print(f"Done: updated {len(changes)} entries.")


if __name__ == "__main__":
    main()

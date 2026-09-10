#!/usr/bin/env python3
"""One-off importer: loads the Hyperview "Operations Alert Actions" runbook
data (the SharePoint list export that used to drive the retired Power App)
into InfraWatch's own hyperview_runbook table.

Usage:
    python3 import_hyperview_runbook.py <path-to-Book1.xlsx> [--db path/to/runbook.db] [--set-by name]

Safe to re-run: rows are upserted by (device_name, alert_type), so importing
the same file twice just refreshes the existing rows instead of duplicating
them. Expects a single sheet with header row:
    Location, Infrastructure Type, Device Type, Device Name, Alert Type,
    Action, Address, Contact Name, Escalation Name
"""
import argparse
import os
import sys
import time

import openpyxl

EXPECTED_HEADER = [
    "Location", "Infrastructure Type", "Device Type", "Device Name", "Alert Type",
    "Action", "Address", "Contact Name", "Escalation Name",
]


def _clean(value):
    """Strips leading/trailing whitespace; a multi-line Action cell's
    internal newlines are preserved."""
    if value is None:
        return ""
    return str(value).strip()


def load_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("Spreadsheet is empty")
    header = [_clean(c) for c in rows[0]]
    if header != EXPECTED_HEADER:
        print(f"WARNING: header doesn't match what's expected.\n  found:    {header}\n  expected: {EXPECTED_HEADER}",
              file=sys.stderr)

    entries = []
    for row in rows[1:]:
        if row is None or all(c is None for c in row):
            continue
        cells = list(row) + [None] * (len(EXPECTED_HEADER) - len(row))
        location, infra_type, device_type, device_name, alert_type, action, address, contact, escalation = cells[:9]
        device_name = _clean(device_name)
        alert_type = _clean(alert_type)
        if not device_name or not alert_type:
            continue
        entries.append({
            "device_name": device_name,
            "alert_type": alert_type,
            "infrastructure_type": _clean(infra_type),
            "device_type": _clean(device_type),
            "location": _clean(location),
            "address": _clean(address),
            "action": _clean(action),
            "primary_name": _clean(contact),
            "primary_phone": "",
            "escalation_name": _clean(escalation),
            "escalation_phone": "",
        })
    return entries


def dedupe(entries):
    """Merge rows that collide on (device_name, alert_type) case-insensitively.
    The source list has a handful of these - same alarm entered twice with a
    slightly different Location spelling/Address completeness. Keep the
    shorter Location (matches the short site-name convention used elsewhere
    in InfraWatch, e.g. "Methodist" not "Methodist Medical Center") and the
    longer Address (usually the one with a ZIP code)."""
    groups = {}
    order = []
    for e in entries:
        key = (e["device_name"].lower(), e["alert_type"].lower())
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(e)

    merged = []
    dupe_count = 0
    for key in order:
        group = groups[key]
        if len(group) == 1:
            merged.append(group[0])
            continue
        dupe_count += len(group) - 1
        base = dict(group[0])
        locations = [g["location"] for g in group if g["location"]]
        addresses = [g["address"] for g in group if g["address"]]
        if locations:
            base["location"] = min(locations, key=len)
        if addresses:
            base["address"] = max(addresses, key=len)
        merged.append(base)
    return merged, dupe_count


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsx_path", help="Path to the runbook export (e.g. Book1.xlsx)")
    parser.add_argument("--db", default=None, help="Path to runbook.db (defaults to RUNBOOK_DB_PATH env var, or ./runbook.db)")
    parser.add_argument("--set-by", default="import", help="Value recorded as updated_by for imported rows (default: 'import')")
    parser.add_argument("--dry-run", action="store_true", help="Parse and report what would be imported, without writing to the db")
    args = parser.parse_args()

    if args.db:
        os.environ["RUNBOOK_DB_PATH"] = args.db

    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import bridge_portal as bp

    entries = load_rows(args.xlsx_path)
    merged, dupe_count = dedupe(entries)
    print(f"Read {len(entries)} row(s) from {args.xlsx_path}")
    if dupe_count:
        print(f"Merged {dupe_count} duplicate device+alert-type row(s) into {len(merged)} unique entries")
    else:
        print(f"{len(merged)} unique device+alert-type entries, no duplicates found")

    if args.dry_run:
        print("--dry-run: not writing to the database. Sample of what would be imported:")
        for e in merged[:5]:
            print(" ", e["device_name"], "/", e["alert_type"], "->", e["location"])
        return

    db_path = os.environ.get("RUNBOOK_DB_PATH", getattr(bp, "RUNBOOK_DB_PATH", "runbook.db"))
    print(f"Writing to {db_path} ...")
    existing = {(e["device_name"].lower(), e["alert_type"].lower()): e for e in bp._hyperview_runbook_entries()}
    created, updated = 0, 0
    for e in merged:
        key = (e["device_name"].lower(), e["alert_type"].lower())
        existing_entry = existing.get(key)
        entry_id = existing_entry["id"] if existing_entry else None
        # preserve any "ignore incomplete" flag someone set by hand in the UI
        ignore_incomplete = bool(existing_entry["ignore_incomplete"]) if existing_entry else False
        bp._save_hyperview_runbook_entry(entry_id, e, args.set_by, ignore_incomplete=ignore_incomplete)
        if entry_id:
            updated += 1
        else:
            created += 1
    print(f"Done: {created} entries created, {updated} entries updated.")


if __name__ == "__main__":
    main()
